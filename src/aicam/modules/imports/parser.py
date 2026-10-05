"""Đọc file nhập đơn `.csv` (UTF-8, có / không BOM) hoặc `.xlsx` (API-50, 02 §6 "Cột mẫu").

Chỉ đọc + kiểm từng ô; không chạm DB. Phân loại NEW / UPDATE / SKIP nằm ở `service.classify`.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from aicam.core.errors import AppError

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
COLUMNS = (
    "platform_order_sn",
    "tracking_number",
    "sku",
    "product_name",
    "variation",
    "quantity",
    "buyer_note",
)
REQUIRED = ("platform_order_sn", "tracking_number", "product_name", "quantity")
MAX_LEN = {"platform_order_sn": 64, "sku": 100, "product_name": 255, "variation": 255, "buyer_note": 500}
EXTENSIONS = {
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def file_invalid(message: str, **details: Any) -> AppError:
    return AppError("FILE_INVALID", message, 422, details)


@dataclass
class Row:
    row: int  # số dòng trong file, dòng tiêu đề = 1
    platform_order_sn: str
    tracking_number: str
    product_name: str
    quantity: int
    sku: str | None = None
    variation: str | None = None
    buyer_note: str | None = None


@dataclass
class RowError:
    row: int
    column: str
    message: str

    def as_dict(self) -> dict[str, Any]:
        return {"row": self.row, "column": self.column, "message": self.message}


@dataclass
class Parsed:
    rows: list[Row] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)


def extension_of(file_name: str) -> str:
    dot = file_name.rfind(".")
    ext = file_name[dot:].lower() if dot >= 0 else ""
    if ext not in EXTENSIONS:
        raise file_invalid("Chỉ nhận file .csv hoặc .xlsx. Dùng file mẫu.", allowed=list(EXTENSIONS))
    return ext


def _cell(value: Any) -> str:
    """Ô Excel kiểu số (mã vận đơn / số lượng gõ thành số) → chuỗi không có `.0`."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _csv_table(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise file_invalid("File CSV phải lưu dạng UTF-8. Dùng file mẫu.", encoding="utf-8") from exc
    first_line = text.split("\n", 1)[0]
    # Excel tiếng Việt hay xuất CSV phân cách `;`.
    delimiter = max((",", ";", "\t"), key=first_line.count)
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _xlsx_table(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # zip hỏng, không phải xlsx…
        raise file_invalid("Không đọc được file Excel. Dùng file mẫu.") from exc
    try:
        sheet = book.worksheets[0]
        table: list[list[str]] = []
        for values in sheet.iter_rows(values_only=True):
            table.append([_cell(v) for v in values])
            if len(table) > MAX_ROWS + 1 + 1000:  # đủ để báo vượt giới hạn, không đọc hết file khổng lồ
                break
        return table
    finally:
        book.close()


def _header(cells: list[str]) -> dict[str, int]:
    index: dict[str, int] = {}
    for i, name in enumerate(cells):
        key = re.sub(r"\s+", "_", name.strip().lower())
        if key in COLUMNS and key not in index:
            index[key] = i
    return index


def parse(content: bytes, ext: str, code_regex: str) -> Parsed:
    """Ném FILE_INVALID cho lỗi cả file (định dạng, thiếu cột, quá 5.000 dòng); lỗi từng ô → `errors`."""
    if len(content) > MAX_BYTES:
        raise file_invalid("File lớn hơn 5 MB.", max_bytes=MAX_BYTES)
    table = _csv_table(content) if ext == ".csv" else _xlsx_table(content)
    if not table:
        raise file_invalid("File trống. Dùng file mẫu.", missing_columns=list(REQUIRED))
    index = _header(table[0])
    missing = [c for c in REQUIRED if c not in index]
    if missing:
        raise file_invalid(
            f"File thiếu cột bắt buộc: {', '.join(missing)}. Dùng file mẫu.", missing_columns=missing
        )
    data = [(n, cells) for n, cells in enumerate(table[1:], start=2) if any(c.strip() for c in cells)]
    if not data:
        raise file_invalid("File không có dòng dữ liệu nào.", max_rows=MAX_ROWS)
    if len(data) > MAX_ROWS:
        raise file_invalid(f"File có hơn {MAX_ROWS} dòng. Chia nhỏ file rồi nhập lại.", max_rows=MAX_ROWS)

    pattern = re.compile(code_regex)
    out = Parsed()
    for n, cells in data:
        values = {c: (cells[i].strip() if i < len(cells) else "") for c, i in index.items()}
        errors: list[RowError] = []
        for column in REQUIRED:
            if not values.get(column):
                errors.append(RowError(n, column, "Bỏ trống"))
        for column, limit in MAX_LEN.items():
            if len(values.get(column, "")) > limit:
                errors.append(RowError(n, column, f"Dài quá {limit} ký tự"))
        tracking = values.get("tracking_number", "").upper()
        if tracking and pattern.fullmatch(tracking) is None:
            errors.append(RowError(n, "tracking_number", "Mã vận đơn không hợp lệ"))
        quantity = 0
        raw_qty = values.get("quantity", "")
        if raw_qty:
            try:
                quantity = int(raw_qty)
            except ValueError:
                quantity = 0
            if quantity <= 0 or quantity > 10000:
                errors.append(RowError(n, "quantity", "Phải là số nguyên từ 1 đến 10000"))
        if errors:
            out.errors.extend(errors)
            continue
        out.rows.append(
            Row(
                row=n,
                platform_order_sn=values["platform_order_sn"],
                tracking_number=tracking,
                product_name=values["product_name"],
                quantity=quantity,
                sku=values.get("sku") or None,
                variation=values.get("variation") or None,
                buyer_note=values.get("buyer_note") or None,
            )
        )
    return out
